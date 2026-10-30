#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
파일럿 문항 인스턴스 생성기
  입력: probe-panel-question-workbook-v1.0.xlsx (템플릿 문항·매개변수 사전), sites.csv, sampling_frame.csv
  출력: instances.jsonl (문항 인스턴스), ledger_template.csv (정답 원장 양식), summary.txt
사용: python generate_instances.py --xlsx ../probe-panel-question-workbook-v1.0.xlsx
"""
import argparse, csv, json, re, itertools, hashlib, collections
from openpyxl import load_workbook

ap = argparse.ArgumentParser()
ap.add_argument("--xlsx", required=True)
ap.add_argument("--sites", default="sites.csv")
ap.add_argument("--frame", default="sampling_frame.csv")
ap.add_argument("--out", default="instances.jsonl")
ap.add_argument("--ledger", default="ledger_template.csv")
ap.add_argument("--ladder", default="L2", help="지명 사다리 기본 단계(L1 광역+기초 / L2 기초 / L3 읍면동 / L4 시설명)")
args = ap.parse_args()

wb = load_workbook(args.xlsx, data_only=True)
ws = wb["템플릿 문항"]
rows = list(ws.iter_rows(min_row=4, values_only=True))
templates = []
for r in rows:
    if not r[0] or not str(r[0]).startswith("Q"):
        continue
    templates.append(dict(qid=r[0], source=r[1], area=r[2], duty=r[3], text=r[4], params=r[5] or "—",
                          qtype=r[6], risk=r[7], variants=r[8], ladder=r[9] or "L2", rule=r[10], doc_type=r[11], law=r[12]))

# 매개변수 사전: 고정 목록(값에 '·'로 나열) vs 지자체별 표집
pws = wb["매개변수 사전"]
fixed = {}
for r in pws.iter_rows(min_row=4, values_only=True):
    if not r[0]: continue
    name, _, values, n = r[0], r[1], r[2] or "", r[3]
    fixed[name] = [v.strip() for v in values.replace(" / ", "·").split("·") if v.strip()]

SITE_PARAMS = {"{동}", "{도서관}", "{축제}", "{명소}", "{시장}", "{공영주차장}", "{보건지소}"}
sites = {r["site_id"]: r for r in csv.DictReader(open(args.sites, encoding="utf-8"))}
frame = collections.defaultdict(list)
for r in csv.DictReader(open(args.frame, encoding="utf-8")):
    frame[(r["site_id"], r["param"])].append(r)

# 템플릿 본문 안의 매개변수 토큰
TOKEN = re.compile(r"\{[^{}]+\}")

def ladder_prefix(site, level):
    name, sido = site["site_name"], site["sido"]
    if level == "L1": return f"{sido} {name}"
    return name  # L2 기본; L3·L4는 토큰 치환으로 구현됨

def values_for(token, site_id, site, tmpl):
    """토큰에 넣을 값 목록과 각 값의 상태(확정/확인 필요/기입 필요/해당 없음)"""
    if token == "{지자체}":
        return [(ladder_prefix(site, args.ladder), "확정")]
    if token in SITE_PARAMS:
        entries = frame.get((site_id, token), [])
        out = []
        for e in entries:
            if e["value"]:
                out.append((e["value"], e["status"]))
            else:
                out.append((f"{token}#{len(out)+1}", "기입 필요"))
        return out
    # 고정 목록: 템플릿의 매개변수 열에 '{품목}: a·b·c' 형태로 적힌 목록을 우선, 없으면 사전
    m = re.search(re.escape(token) + r"\s*:\s*([^/\n]+)", tmpl["params"])
    if m and "종" not in m.group(1) and "위와 동일" not in m.group(1) and "표집" not in m.group(1):
        vals = [v.strip() for v in m.group(1).replace("(", "·").replace(")", "").split("·") if v.strip()]
        vals = [v for v in vals if not v.endswith("대상")]
        if vals: return [(v, "확정") for v in vals]
    key = token
    if key not in fixed:
        # {차종}×{구역} 같은 복합 키 처리
        for k in fixed:
            if token in k: key = k; break
    vals = fixed.get(key, [])
    if token == "{용량}":
        vals = ["2L", "3L", "5L"] if "음식물" in tmpl["text"] else ["10L", "20L", "50L", "75L"]
    return [(v, "확정") for v in vals] or [(token, "기입 필요")]

instances = []
skipped = collections.Counter()
for site_id, site in sites.items():
    for t in templates:
        tokens = TOKEN.findall(t["text"])
        tokens = list(dict.fromkeys(tokens))
        choices = []
        for tok in tokens:
            vals = values_for(tok, site_id, site, t)
            if not vals:
                skipped[(site_id, tok)] += 1
            choices.append([(tok, v, st) for v, st in vals] or [(tok, None, "해당 없음")])
        for combo in itertools.product(*choices):
            if any(v is None for _, v, _ in combo):
                skipped[(site_id, "해당 없음:" + t["qid"])] += 1
                continue
            text = t["text"]
            params = {}
            status = "확정"
            for tok, v, st in combo:
                text = text.replace(tok, v)
                if tok != "{지자체}": params[tok] = v
                if st != "확정": status = st if status == "확정" or st == "기입 필요" else status
            ladder = t["ladder"]
            key = f"{site_id}|{t['qid']}|{json.dumps(params, ensure_ascii=False, sort_keys=True)}"
            iid = hashlib.sha1(key.encode()).hexdigest()[:10]
            instances.append(dict(
                instance_id=iid, site_id=site_id, site_name=site["site_name"], qid=t["qid"], area=t["area"], duty=t["duty"],
                source=t["source"], qtype=t["qtype"], risk=t["risk"], ladder=ladder, params=params,
                question=text, rule=t["rule"], doc_type=t["doc_type"], law=t["law"], status=status,
                official_domains=site["official_domains"].split(";")))

with open(args.out, "w", encoding="utf-8") as f:
    for i in instances:
        f.write(json.dumps(i, ensure_ascii=False) + "\n")

# 정답 원장 양식: 고위험 먼저, 기입 필요 항목은 뒤로
order = {"고": 0, "중": 1, "저": 2}
instances.sort(key=lambda i: (i["status"] != "확정", order[i["risk"]], i["site_id"], i["qid"]))
with open(args.ledger, "w", encoding="utf-8", newline="") as f:
    w = csv.writer(f)
    w.writerow(["instance_id", "site", "qid", "risk", "qtype", "question", "status",
                "official_url_1", "official_url_2(고위험 필수)", "answer_key(정답 요지)", "numbers(숫자·날짜)",
                "supply_gap(공식 문서 없음=Y)", "ledger_author", "second_author(고위험)", "checked_date", "note"])
    for i in instances:
        w.writerow([i["instance_id"], i["site_name"], i["qid"], i["risk"], i["qtype"], i["question"], i["status"]] + [""] * 9)

c = collections.Counter((i["site_id"], i["status"]) for i in instances)
with open("summary.txt", "w", encoding="utf-8") as f:
    f.write(f"템플릿 {len(templates)} · 인스턴스 {len(instances)}\n")
    for s in sites:
        n = sum(v for (sid, _), v in c.items() if sid == s)
        f.write(f"{sites[s]['site_name']}: {n}  (확정 {c[(s,'확정')]}, 확인 필요 {c[(s,'확인 필요')]}, 기입 필요 {c[(s,'기입 필요')]})\n")
    f.write("제외(해당 없음·값 없음):\n")
    for k, v in sorted(skipped.items()):
        f.write(f"  {k}: {v}\n")
print(open("summary.txt", encoding="utf-8").read())
